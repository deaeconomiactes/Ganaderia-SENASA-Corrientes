"""Read an internal workbook without printing cell values to the console.

The JavaScript pipeline owns validation and output shaping. This helper only
bridges XLSX -> JSON rows using openpyxl, which is available in the local DEA
runtime. It intentionally emits one JSON document on stdout for the caller.
"""

from __future__ import annotations

import json
import sys
from datetime import date, datetime, time
from pathlib import Path


def json_value(value):
    if isinstance(value, (datetime, date, time)):
        return value.isoformat()
    return value


def main() -> int:
    # Windows can default stdout to a legacy code page. The JSON bridge must
    # always emit UTF-8 so locality names never break the Node.js pipeline.
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(encoding="utf-8")
    if hasattr(sys.stderr, "reconfigure"):
        sys.stderr.reconfigure(encoding="utf-8")
    if len(sys.argv) != 2:
        print("Uso: read-xlsx-json.py RUTA_XLSX", file=sys.stderr)
        return 2
    source = Path(sys.argv[1])
    try:
        import openpyxl
    except ImportError:
        print("Falta openpyxl. Instale openpyxl en el entorno local para leer XLSX.", file=sys.stderr)
        return 3
    try:
        workbook = openpyxl.load_workbook(source, read_only=True, data_only=True)
        expected_sheets = ("Agricola", "Ganadero", "Mixto")
        source_sheets = [name for name in expected_sheets if name in workbook.sheetnames]
        if source_sheets and len(source_sheets) != len(expected_sheets):
            print("Faltan hojas requeridas: Agricola, Ganadero y Mixto.", file=sys.stderr)
            return 4
        if not source_sheets:
            source_sheets = [workbook.sheetnames[0]]
        payload = []
        expected_headers = None
        for sheet_name in source_sheets:
            sheet = workbook[sheet_name]
            rows = sheet.iter_rows(values_only=True)
            headers = next(rows, None)
            if not headers:
                print(f"La hoja {sheet_name} no contiene una fila de cabeceras.", file=sys.stderr)
                return 4
            clean_headers = [str(value or "").strip() for value in headers]
            if expected_headers is not None and clean_headers != expected_headers:
                print(f"La estructura de la hoja {sheet_name} difiere de las demás.", file=sys.stderr)
                return 4
            expected_headers = clean_headers
            for values in rows:
                row = {header: json_value(value) for header, value in zip(clean_headers, values) if header}
                if any(value not in (None, "") for value in row.values()):
                    row["__sourceSheet"] = sheet_name
                    payload.append(row)
        json.dump(payload, sys.stdout, ensure_ascii=False, separators=(",", ":"))
        return 0
    except Exception as error:  # pragma: no cover - surfaced as a concise CLI error
        print(f"No se pudo leer el XLSX: {error}", file=sys.stderr)
        return 5


if __name__ == "__main__":
    raise SystemExit(main())
