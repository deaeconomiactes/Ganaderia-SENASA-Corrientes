#!/usr/bin/env python3
"""Genera un reporte de titulares con tres o mas RENSPA unicos.

El archivo fuente se abre solo para lectura y nunca se modifica.
"""

from __future__ import annotations

import argparse
import re
import sys
import unicodedata
from collections import defaultdict
from datetime import datetime
from pathlib import Path
from typing import Any, Iterable

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


CUIL_CANDIDATES = (
    "CUIL",
    "CUIT",
    "CUIT/CUIL",
    "CUIT_CUIL",
    "DOCUMENTO",
    "NRO_DOCUMENTO",
    "NUMERO_DOCUMENTO",
    "TITULAR_DOCUMENTO",
    "DOC_PRODUCTOR",
    "IDENTIFICADOR",
)
RENSPA_CANDIDATES = (
    "RENSPA",
    "UP_RENSPA",
    "REGISTRO_RENSPA",
    "UNIDAD_PRODUCTIVA",
    "ID_RENSPA",
)

SUMMARY_DIMENSIONS = {
    "Departamentos": ("DEPARTAMENTO", "DEPTO"),
    "Municipios": ("MUNICIPIO", "MUNI"),
    "Oficinas_Locales": ("OFICINA_LOCAL", "OFICINA LOCAL"),
}
SUMMARY_STOCKS = {
    "Existencias_Totales": (
        "EXISTENCIAS_TOTALES",
        "EXISTENCIA_TOTAL",
        "TOTAL_EXISTENCIAS",
        "TOTAL_GANADERO",
        "TOTAL_GANADO",
    ),
    "Bovinos": ("BOVINOS",),
    "Bubalinos": ("BUBALINOS",),
    "Ovinos": ("OVINOS",),
    "Caprinos": ("CAPRINOS",),
    "Porcinos": ("PORCINOS",),
    "Equinos": ("EQUINOS",),
}


def normalized_header(value: Any) -> str:
    """Normaliza encabezados para compararlos sin depender de puntuacion."""
    text = "" if value is None else str(value).strip()
    text = "".join(
        char for char in unicodedata.normalize("NFKD", text) if not unicodedata.combining(char)
    )
    return re.sub(r"[^A-Z0-9]+", "", text.upper())


def text_value(value: Any) -> str:
    """Convierte identificadores a texto sin agregar el sufijo .0."""
    if value is None:
        return ""
    if isinstance(value, bool):
        return str(value)
    if isinstance(value, int):
        return str(value)
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def normalize_cuil(value: Any) -> str:
    return "".join(re.findall(r"\d", text_value(value)))


def normalize_renspa(value: Any) -> str:
    return "".join(re.findall(r"[A-Za-z0-9]", text_value(value))).upper()


def find_exact(headers: list[Any], candidates: Iterable[str]) -> int | None:
    aliases = {normalized_header(candidate) for candidate in candidates}
    for index, header in enumerate(headers):
        if normalized_header(header) in aliases:
            return index
    return None


def documentary_score(header: Any) -> int:
    key = normalized_header(header)
    if not key:
        return -1
    score = 0
    if "CUIL" in key:
        score += 120
    if "CUIT" in key:
        score += 110
    if "DOCUMENTO" in key:
        score += 80
    if key.startswith("DOC") or "IDENTIFICADOR" in key:
        score += 45
    if "TITULAR" in key:
        score += 25
    if "PRODUCTOR" in key:
        score += 20
    if "TELEF" in key or "MAIL" in key:
        score -= 200
    return score


def detect_columns(headers: list[Any]) -> tuple[int | None, bool, int | None]:
    # Primero prioriza identificadores tributarios explicitos. Si no existe una
    # coincidencia exacta, compara todas las alternativas documentales: una
    # variante como "CUIT/L TITULAR UP" es mas especifica que NRO_DOCUMENTO.
    cuil_index = find_exact(headers, ("CUIL", "CUIT", "CUIT/CUIL", "CUIT_CUIL"))
    used_fallback = False
    if cuil_index is None:
        ranked = sorted(
            ((documentary_score(header), index) for index, header in enumerate(headers)),
            reverse=True,
        )
        if ranked and ranked[0][0] > 0:
            cuil_index = ranked[0][1]
            used_fallback = True
    renspa_index = find_exact(headers, RENSPA_CANDIDATES)
    return cuil_index, used_fallback, renspa_index


def detect_sheet_and_header(source: Path) -> tuple[str, int, list[Any], int, bool, int]:
    """Busca una fila de encabezados util en todas las hojas (primeras 25 filas)."""
    workbook = load_workbook(source, read_only=True, data_only=True)
    options: list[tuple[int, int, str, int, list[Any], int, bool, int]] = []
    try:
        for sheet in workbook.worksheets:
            for row_number, row in enumerate(
                sheet.iter_rows(min_row=1, max_row=min(sheet.max_row, 25), values_only=True), start=1
            ):
                headers = list(row)
                cuil_index, used_fallback, renspa_index = detect_columns(headers)
                score = (100 if cuil_index is not None else 0) + (100 if renspa_index is not None else 0)
                score += sum(value not in (None, "") for value in headers)
                options.append(
                    (
                        score,
                        sheet.max_row,
                        sheet.title,
                        row_number,
                        headers,
                        -1 if cuil_index is None else cuil_index,
                        used_fallback,
                        -1 if renspa_index is None else renspa_index,
                    )
                )
        if not options:
            raise ValueError("El archivo no contiene hojas legibles.")
        best = max(options, key=lambda item: (item[0], item[1], -item[3]))
        _, _, sheet_name, row_number, headers, cuil_index, used_fallback, renspa_index = best
        return sheet_name, row_number, headers, cuil_index, used_fallback, renspa_index
    finally:
        workbook.close()


def number_value(value: Any) -> float:
    if value in (None, ""):
        return 0.0
    if isinstance(value, bool):
        return 0.0
    if isinstance(value, (int, float)):
        return float(value)
    candidate = str(value).strip().replace(" ", "")
    if not candidate:
        return 0.0
    # Admite representaciones usuales 1.234,56 y 1234.56.
    if "," in candidate and "." in candidate:
        candidate = candidate.replace(".", "").replace(",", ".")
    elif "," in candidate:
        candidate = candidate.replace(",", ".")
    try:
        return float(candidate)
    except ValueError:
        return 0.0


def unique_join(values: Iterable[Any]) -> str:
    cleaned = {text_value(value) for value in values if text_value(value)}
    return ", ".join(sorted(cleaned, key=lambda item: item.casefold()))


def compatible_index(headers: list[Any], aliases: Iterable[str]) -> int | None:
    alias_keys = {normalized_header(alias) for alias in aliases}
    for index, header in enumerate(headers):
        if normalized_header(header) in alias_keys:
            return index
    return None


def relevant_detail_index(header: Any) -> bool:
    key = normalized_header(header)
    if not key:
        return False
    # Los datos de contacto no forman parte del alcance de ubicacion/stock y
    # agregarian exposicion innecesaria de informacion personal.
    if any(token in key for token in ("TELEF", "CELULAR", "MAIL", "EMAIL", "CORREO")):
        return False
    relevant_tokens = (
        "PROVINCIA",
        "DEPTO",
        "DEPARTAMENTO",
        "MUNI",
        "MUNICIPIO",
        "OFICINA",
        "LOCALIDAD",
        "ESTABLECIMIENTO",
        "TITULAR",
        "RAZONSOCIAL",
        "NOMBRE",
        "LATITUD",
        "LONGITUD",
        "BOV",
        "BUB",
        "EQUIN",
        "CABALLO",
        "YEGUA",
        "ASNO",
        "BURRO",
        "MULA",
        "PORC",
        "CERDA",
        "LECHON",
        "CAPR",
        "CABRA",
        "CHIVO",
        "OVIN",
        "OVEJA",
        "CARNERO",
        "BORREGO",
        "CORDER",
        "CAMEL",
        "EXISTENCIA",
        "STOCK",
        "TOTAL",
        "VACAS",
        "TOROS",
        "BUEYES",
        "NOVILL",
        "VAQUILL",
        "TERNER",
        "POTRILL",
        "PADRILL",
        "CACHORR",
        "CAPONES",
        "MUFLON",
    )
    return any(token in key for token in relevant_tokens)


def unique_headers(headers: list[Any]) -> list[str]:
    used: dict[str, int] = {}
    output: list[str] = []
    for position, header in enumerate(headers, start=1):
        base = text_value(header) or f"Columna_{position}"
        count = used.get(base.casefold(), 0) + 1
        used[base.casefold()] = count
        output.append(base if count == 1 else f"{base}_{count}")
    return output


def style_sheet(sheet, text_columns: set[int], highlighted_column: int | None = None) -> None:
    header_fill = PatternFill("solid", fgColor="1F4E78")
    highlight_fill = PatternFill("solid", fgColor="D9EAF7")
    white_bold = Font(name="Arial", size=10, bold=True, color="FFFFFF")
    body_font = Font(name="Arial", size=10, color="1F1F1F")

    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    sheet.sheet_view.showGridLines = False
    sheet.row_dimensions[1].height = 30

    for cell in sheet[1]:
        cell.fill = header_fill
        cell.font = white_bold
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.font = body_font
            cell.alignment = Alignment(vertical="top")
            if cell.column in text_columns:
                cell.number_format = "@"
            elif isinstance(cell.value, (int, float)) and not isinstance(cell.value, bool):
                cell.number_format = "#,##0"
        if highlighted_column:
            row[highlighted_column - 1].fill = highlight_fill
            row[highlighted_column - 1].font = Font(name="Arial", size=10, bold=True, color="1F1F1F")

    for column_cells in sheet.iter_cols():
        letter = get_column_letter(column_cells[0].column)
        observed = [len(text_value(cell.value)) for cell in column_cells[: min(len(column_cells), 500)]]
        width = min(max(max(observed, default=8) + 2, 10), 45)
        header_key = normalized_header(column_cells[0].value)
        if "LISTADO" in header_key:
            width = 45
        elif header_key in {"DEPARTAMENTOS", "MUNICIPIOS", "OFICINASLOCALES"}:
            width = min(max(width, 22), 35)
        sheet.column_dimensions[letter].width = width


def generate_report(source: Path, output: Path) -> dict[str, Any]:
    sheet_name, header_row, headers, cuil_index, fallback, renspa_index = detect_sheet_and_header(source)
    if cuil_index < 0:
        raise ValueError("No se encontro una columna de CUIL ni una alternativa documental compatible.")
    if renspa_index < 0:
        raise ValueError("No se encontro una columna de RENSPA compatible.")

    workbook = load_workbook(source, read_only=True, data_only=True)
    try:
        sheet = workbook[sheet_name]
        rows = [
            tuple(row)
            for row in sheet.iter_rows(min_row=header_row + 1, values_only=True)
            if any(value not in (None, "") for value in row)
        ]
    finally:
        workbook.close()

    normalized_rows: list[tuple[tuple[Any, ...], str, str]] = []
    renspas_by_cuil: dict[str, set[str]] = defaultdict(set)
    for row in rows:
        cuil = normalize_cuil(row[cuil_index] if cuil_index < len(row) else None)
        renspa = normalize_renspa(row[renspa_index] if renspa_index < len(row) else None)
        normalized_rows.append((row, cuil, renspa))
        if cuil and renspa:
            renspas_by_cuil[cuil].add(renspa)

    qualifying = {cuil for cuil, renspas in renspas_by_cuil.items() if len(renspas) >= 3}
    detail_records = [record for record in normalized_rows if record[1] in qualifying]

    dimension_indexes = {
        output_name: compatible_index(headers, aliases)
        for output_name, aliases in SUMMARY_DIMENSIONS.items()
    }
    stock_indexes = {
        output_name: compatible_index(headers, aliases)
        for output_name, aliases in SUMMARY_STOCKS.items()
    }
    stock_indexes = {name: index for name, index in stock_indexes.items() if index is not None}

    summary_headers = [
        "CUIL",
        "Cantidad_RENSPA",
        "Cantidad_Registros",
        "RENSPA_Listado",
        "Departamentos",
        "Municipios",
        "Oficinas_Locales",
        *stock_indexes.keys(),
    ]
    records_by_cuil: dict[str, list[tuple[Any, ...]]] = defaultdict(list)
    for row, cuil, _ in detail_records:
        records_by_cuil[cuil].append(row)

    summary_rows: list[list[Any]] = []
    for cuil in qualifying:
        cuil_rows = records_by_cuil[cuil]
        values: list[Any] = [
            cuil,
            len(renspas_by_cuil[cuil]),
            len(cuil_rows),
            ", ".join(sorted(renspas_by_cuil[cuil])),
        ]
        for name in SUMMARY_DIMENSIONS:
            index = dimension_indexes[name]
            values.append(unique_join(row[index] for row in cuil_rows) if index is not None else "")
        for index in stock_indexes.values():
            total = sum(number_value(row[index]) for row in cuil_rows)
            values.append(int(total) if total.is_integer() else total)
        summary_rows.append(values)
    summary_rows.sort(key=lambda row: (-row[1], row[0]))

    candidate_detail_indexes = [
        index
        for index, header in enumerate(headers)
        if index not in {cuil_index, renspa_index} and relevant_detail_index(header)
    ]
    nonempty_detail_indexes = [
        index
        for index in candidate_detail_indexes
        if any(index < len(row) and row[index] not in (None, "") for row, _, _ in detail_records)
    ]
    detail_headers = ["CUIL", "RENSPA"] + unique_headers([headers[index] for index in nonempty_detail_indexes])
    detail_rows: list[list[Any]] = []
    for row, cuil, renspa in detail_records:
        detail_rows.append(
            [cuil, renspa]
            + [row[index] if index < len(row) else None for index in nonempty_detail_indexes]
        )

    def detail_sort_key(values: list[Any]) -> tuple[str, str, str, str]:
        lookup = {normalized_header(header): values[index] for index, header in enumerate(detail_headers)}
        department = lookup.get("DEPTO", lookup.get("DEPARTAMENTO", ""))
        municipality = lookup.get("MUNI", lookup.get("MUNICIPIO", ""))
        return values[0], values[1], text_value(department).casefold(), text_value(municipality).casefold()

    detail_rows.sort(key=detail_sort_key)

    output.parent.mkdir(parents=True, exist_ok=True)
    report = Workbook()
    summary_sheet = report.active
    summary_sheet.title = "Resumen_CUIL"
    detail_sheet = report.create_sheet("Detalle_RENSPA")
    metadata_sheet = report.create_sheet("Metadatos")

    summary_sheet.append(summary_headers)
    for row in summary_rows:
        summary_sheet.append(row)
    detail_sheet.append(detail_headers)
    for row in detail_rows:
        detail_sheet.append(row)

    metadata = [
        ("Campo", "Valor"),
        ("Archivo fuente", source.name),
        ("Hoja fuente", sheet_name),
        ("Fila de encabezados", header_row),
        ("Fecha y hora de generacion", datetime.now().astimezone().isoformat(timespec="seconds")),
        ("Columna usada como CUIL", text_value(headers[cuil_index])),
        ("Columna usada como RENSPA", text_value(headers[renspa_index])),
        ("Total de filas leidas", len(rows)),
        ("Total de CUIL unicos", len(renspas_by_cuil)),
        ("Total de CUIL con 3 RENSPA o mas", len(qualifying)),
        ("Total de RENSPA involucrados", sum(len(renspas_by_cuil[cuil]) for cuil in qualifying)),
        (
            "Criterio aplicado",
            "CUIL normalizado a digitos y RENSPA normalizado a letras/numeros; se incluyen CUIL con 3 o mas RENSPA unicos.",
        ),
    ]
    for row in metadata:
        metadata_sheet.append(row)

    style_sheet(summary_sheet, {1}, highlighted_column=2)
    style_sheet(detail_sheet, {1, 2})
    style_sheet(metadata_sheet, set())
    for row in summary_sheet.iter_rows(min_row=2, min_col=4, max_col=7):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    for row_number in range(2, summary_sheet.max_row + 1):
        summary_sheet.row_dimensions[row_number].height = 45
    metadata_sheet.column_dimensions["A"].width = 34
    metadata_sheet.column_dimensions["B"].width = 75
    for cell in metadata_sheet["B"]:
        cell.alignment = Alignment(vertical="top", wrap_text=True)
    metadata_sheet.row_dimensions[12].height = 34

    report.save(output)
    report.close()
    return {
        "sheet": sheet_name,
        "header_row": header_row,
        "cuil_column": text_value(headers[cuil_index]),
        "renspa_column": text_value(headers[renspa_index]),
        "cuil_fallback": fallback,
        "rows_read": len(rows),
        "unique_cuil": len(renspas_by_cuil),
        "qualifying_cuil": len(qualifying),
        "involved_renspa": sum(len(renspas_by_cuil[cuil]) for cuil in qualifying),
        "detail_rows": len(detail_rows),
    }


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description="Genera un XLSX con CUIL que tienen al menos 3 RENSPA unicos."
    )
    parser.add_argument("fuente", type=Path, help="Ruta al archivo XLSX fuente")
    parser.add_argument("salida", nargs="?", type=Path, help="Ruta opcional al XLSX de salida")
    return parser.parse_args()


def main() -> int:
    args = parse_args()
    source = args.fuente.expanduser().resolve()
    if not source.is_file():
        print(f"ERROR: no existe el archivo fuente: {source}", file=sys.stderr)
        return 2
    if source.suffix.lower() != ".xlsx":
        print("ERROR: el archivo fuente debe ser XLSX.", file=sys.stderr)
        return 2
    output = (
        args.salida.expanduser().resolve()
        if args.salida
        else source.with_name("reporte_cuil_3_o_mas_renspa.xlsx")
    )
    if output == source:
        print("ERROR: la salida no puede sobrescribir el archivo fuente.", file=sys.stderr)
        return 2

    try:
        metrics = generate_report(source, output)
    except (ValueError, OSError) as error:
        print(f"ERROR: {error}", file=sys.stderr)
        return 1

    if metrics["cuil_fallback"]:
        print(
            "AVISO: no se encontro un encabezado CUIL exacto; "
            f"se uso la mejor columna documental disponible: {metrics['cuil_column']}."
        )
    if metrics["qualifying_cuil"] == 0:
        print("AVISO: no se encontraron CUIL con 3 RENSPA unicos o mas.")
    print(f"Reporte creado: {output}")
    print(f"Columna CUIL: {metrics['cuil_column']}")
    print(f"Columna RENSPA: {metrics['renspa_column']}")
    print(f"Filas leidas: {metrics['rows_read']:,}")
    print(f"CUIL unicos: {metrics['unique_cuil']:,}")
    print(f"CUIL con 3 RENSPA o mas: {metrics['qualifying_cuil']:,}")
    print(f"RENSPA involucrados: {metrics['involved_renspa']:,}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
